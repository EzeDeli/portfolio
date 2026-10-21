"""Carga automática de las planillas de ventas de cada sucursal.

Busca Excel nuevos en entrada/, valida y limpia cada fila, carga las buenas en la base,
separa las que tienen errores con el motivo, mueve el archivo a procesados/ y manda un resumen por mail.
Está pensado para correr solo, cada 15 minutos, con el Programador de tareas de Windows o cron.

Uso:  python proceso.py
"""

import csv
import datetime as dt
import hashlib
import logging
import os
import re
import shutil
import sqlite3
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

import aviso
from validaciones import ErrorDeFila, limpiar_fila, mapear_encabezados

CARPETA = Path(__file__).resolve().parent
ENTRADA = CARPETA / "entrada"
PROCESADOS = CARPETA / "procesados"
SALIDA = CARPETA / "salida"
BASE = SALIDA / "ventas.db"

log = logging.getLogger("carga")


def conectar():
    SALIDA.mkdir(exist_ok=True)
    base = sqlite3.connect(BASE)
    base.executescript((CARPETA / "esquema.sql").read_text(encoding="utf-8"))
    with open(CARPETA / "productos.csv", encoding="utf-8") as f:
        productos = [(r["codigo"], r["nombre"], float(r["precio_lista"])) for r in csv.DictReader(f, delimiter=";")]
    base.executemany("INSERT OR REPLACE INTO productos VALUES (?, ?, ?)", productos)
    base.commit()
    return base


def huella(ruta):
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def sucursal_de(ruta):
    """ventas_centro_2026-10-20.xlsx -> 'centro'"""
    coincide = re.match(r"ventas_([a-z]+)_", ruta.name.lower())
    return coincide.group(1) if coincide else ruta.stem


def mostrar(valor):
    """Cómo se ve el valor original en el reporte de rechazos."""
    if valor is None:
        return ""
    if isinstance(valor, (dt.date, dt.datetime)):
        return f"{valor:%d/%m/%Y}"
    return str(valor)


def procesar_archivo(base, ruta, catalogo):
    """Carga un Excel. Devuelve el resumen y la lista de filas rechazadas."""
    resumen = {"archivo": ruta.name, "sucursal": sucursal_de(ruta), "leidas": 0,
               "cargadas": 0, "rechazadas": 0, "total": 0.0, "estado": "ok"}
    rechazos = []

    if base.execute("SELECT 1 FROM archivos WHERE huella = ?", (huella(ruta),)).fetchone():
        resumen["estado"] = "repetido: ya se había cargado este mismo archivo"
        return resumen, rechazos

    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:                                           # en modo lectura el archivo queda abierto hasta cerrarlo
        filas = list(libro.active.iter_rows(values_only=True))
    finally:
        libro.close()
    columnas = mapear_encabezados(filas[0])
    filas = filas[1:]
    ahora = dt.datetime.now().isoformat(timespec="seconds")

    with base:                                     # todo el archivo entra junto, o no entra nada
        for numero, fila in enumerate(filas, start=2):
            if all(v is None or str(v).strip() == "" for v in fila):
                continue                           # filas vacías: se ignoran sin contarlas
            resumen["leidas"] += 1
            try:
                venta = limpiar_fila(fila, columnas, catalogo)
                base.execute(
                    "INSERT INTO ventas VALUES (:id_venta, :sucursal, :fecha, :producto, :cantidad,"
                    " :precio, :total, :archivo, :cargado_en)",
                    {**venta, "fecha": venta["fecha"].isoformat(), "sucursal": resumen["sucursal"],
                     "archivo": ruta.name, "cargado_en": ahora})
            except ErrorDeFila as e:
                motivo = str(e)
            except sqlite3.IntegrityError:
                motivo = f"Venta repetida (ya estaba cargada): {venta['id_venta']}"
            else:
                resumen["cargadas"] += 1
                resumen["total"] += venta["total"]
                continue
            resumen["rechazadas"] += 1
            rechazos.append({"archivo": ruta.name, "fila": numero, "motivo": motivo,
                             "datos": " | ".join(mostrar(v) for v in fila)})

        base.execute("INSERT INTO archivos VALUES (?, ?, ?, ?, ?)",
                     (huella(ruta), ruta.name, ahora, resumen["cargadas"], resumen["rechazadas"]))
    return resumen, rechazos


def guardar_rechazos(rechazos):
    if not rechazos:
        return None
    ruta = SALIDA / f"rechazos_{dt.datetime.now():%Y-%m-%d_%H%M%S}.csv"
    with open(ruta, "w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig: Excel lo abre bien
        w = csv.DictWriter(f, fieldnames=["archivo", "fila", "motivo", "datos"], delimiter=";")
        w.writeheader()
        w.writerows(rechazos)
    return ruta


def cargar_env():
    """Lee .env si existe (sin librerías extra). Las variables ya definidas tienen prioridad."""
    archivo = CARPETA / ".env"
    if archivo.exists():
        for linea in archivo.read_text(encoding="utf-8").splitlines():
            clave, signo, valor = linea.partition("=")
            if signo and not linea.lstrip().startswith("#"):
                os.environ.setdefault(clave.strip(), valor.strip())


def main():
    cargar_env()
    for carpeta in (ENTRADA, PROCESADOS, SALIDA):
        carpeta.mkdir(exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                        handlers=[logging.FileHandler(SALIDA / "proceso.log", encoding="utf-8"),
                                  logging.StreamHandler()])

    archivos = sorted(ENTRADA.glob("*.xlsx"))
    if not archivos:
        log.info("No hay archivos nuevos")
        return

    base = conectar()
    catalogo = {c for (c,) in base.execute("SELECT codigo FROM productos")}
    resumenes, rechazos = [], []

    for ruta in archivos:
        try:
            resumen, rechazadas = procesar_archivo(base, ruta, catalogo)
        except Exception as e:                       # un archivo roto no frena a los demás
            log.exception("No se pudo procesar %s", ruta.name)
            resumen = {"archivo": ruta.name, "sucursal": sucursal_de(ruta), "leidas": 0, "cargadas": 0,
                       "rechazadas": 0, "total": 0.0, "estado": f"error: {e}"}
            rechazadas = []
        resumenes.append(resumen)
        rechazos.extend(rechazadas)
        shutil.move(str(ruta), PROCESADOS / ruta.name)
        log.info("%s: %s cargadas, %s rechazadas (%s)", ruta.name, resumen["cargadas"],
                 resumen["rechazadas"], resumen["estado"])

    archivo_rechazos = guardar_rechazos(rechazos)
    motivos = Counter(r["motivo"].split(":")[0] for r in rechazos)
    aviso.enviar_resumen(resumenes, rechazos, motivos, archivo_rechazos, SALIDA)
    base.close()


if __name__ == "__main__":
    main()
